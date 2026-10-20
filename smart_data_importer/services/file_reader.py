import base64
import csv
import datetime
import hashlib
import io
import logging
import unicodedata
import warnings
from collections import OrderedDict

import chardet
import pandas as pd

from odoo import _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Sheet kinds
SHEET_DATA = 'data'
SHEET_INSTRUCTIONS = 'instructions'
SHEET_VALUES = 'values'
SHEET_EMPTY = 'empty'

# Header colour legend levels (read from the "Instrucciones" sheet)
LEVEL_REQUIRED = 'required'
LEVEL_CONDITIONAL = 'conditional'
LEVEL_OPTIONAL = 'optional'

_LEGEND_WORDS = {
    'obligatorio': LEVEL_REQUIRED,
    'requerido': LEVEL_REQUIRED,
    'required': LEVEL_REQUIRED,
    'condicional': LEVEL_CONDITIONAL,
    'conditional': LEVEL_CONDITIONAL,
    'opcional': LEVEL_OPTIONAL,
    'optional': LEVEL_OPTIONAL,
}

_EXCEL_ERRORS = {'#N/A', '#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#NULL!', '#NUM!'}
_MARKER_VALUES = {'1', 'x', 'X'}

# Small in-process cache: the same file is read by the analysis, the preview,
# the duplicate counter and every import chunk.
_CACHE = OrderedDict()
_CACHE_SIZE = 6


class FileReader:
    """Reads CSV and Excel files into clean, all-text DataFrames.

    Excel workbooks may hold several sheets. Each sheet is classified as
    instructions, "possible values" (lookup lists), data, or empty, and the
    caller picks which data sheet to import.

    Every value is returned as a string (or None when empty) so identifiers,
    postal codes and phone numbers keep their leading zeros and never turn
    into floats such as '3001234567.0'.
    """

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    @staticmethod
    def normalize_header(header):
        """Shared normalization: trim, lowercase, remove diacritics."""
        if header is None:
            return ''
        header = str(header).strip().lower()
        header = ''.join(c for c in unicodedata.normalize('NFD', header)
                         if unicodedata.category(c) != 'Mn')
        return header

    @classmethod
    def read(cls, file_base64, filename, sheet_name=None):
        """Read one sheet (or the CSV) and return a DataFrame plus metadata.

        :param sheet_name: Excel sheet to read. When empty, the first visible
            data sheet is used.
        :return: dict with keys df, encoding, delimiter, header_row_index,
            total_rows, sheet_name, column_meta, dropped_rows, sheets
        """
        workbook = cls._parse(file_base64, filename)
        sheets = workbook['sheets']
        if not sheets:
            raise UserError(_("The file does not contain any readable sheet."))

        chosen = None
        if sheet_name:
            chosen = next((s for s in sheets if s['name'] == sheet_name), None)
            if chosen is None:
                raise UserError(_("The sheet '%s' does not exist in this file.") % sheet_name)
        else:
            chosen = cls._default_sheet(sheets)

        return {
            'df': chosen['df'].copy(),
            'encoding': workbook['encoding'],
            'delimiter': workbook['delimiter'],
            'header_row_index': chosen['header_row_index'],
            'total_rows': len(chosen['df']),
            'sheet_name': chosen['name'],
            'column_meta': dict(chosen['column_meta']),
            'dropped_rows': chosen['dropped_rows'],
            'sheets': [cls._sheet_summary(s) for s in sheets],
        }

    @classmethod
    def list_sheets(cls, file_base64, filename):
        """Return a light summary of every sheet (no DataFrames)."""
        workbook = cls._parse(file_base64, filename)
        return [cls._sheet_summary(s) for s in workbook['sheets']]

    # ------------------------------------------------------------------
    # Parsing
    # ------------------------------------------------------------------

    @classmethod
    def _parse(cls, file_base64, filename):
        if not file_base64:
            raise UserError(_("No file content found."))
        if not filename:
            raise UserError(_("The file has no name, so its format cannot be detected."))

        content = base64.b64decode(file_base64)
        key = (hashlib.sha1(content).hexdigest(), filename.lower())
        if key in _CACHE:
            _CACHE.move_to_end(key)
            return _CACHE[key]

        lower = filename.lower()
        if lower.endswith('.csv'):
            result = cls._parse_csv(content)
        elif lower.endswith(('.xlsx', '.xlsm')):
            result = cls._parse_xlsx(content)
        elif lower.endswith('.xls'):
            result = cls._parse_xls_legacy(content)
        else:
            raise UserError(_("Unsupported file format. Use CSV or Excel (.xlsx, .xls)."))

        _CACHE[key] = result
        if len(_CACHE) > _CACHE_SIZE:
            _CACHE.popitem(last=False)
        return result

    @classmethod
    def _parse_csv(cls, content):
        detected = chardet.detect(content[:10000])
        encoding = detected.get('encoding') or 'utf-8'
        try:
            text = content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            encoding = 'latin-1'
            text = content.decode(encoding)
        text = text.lstrip('﻿')

        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|').delimiter
        except csv.Error:
            delimiter = ','

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        grid = [[cls._to_text(v) for v in row] for row in reader]
        # CSV line numbers are 1-based, same as spreadsheet rows
        sheet = cls._build_sheet('CSV', grid, visible=True, fills=None, legend=None, force_data=True)
        return {'encoding': encoding, 'delimiter': delimiter, 'sheets': [sheet]}

    @classmethod
    def _parse_xlsx(cls, content):
        import openpyxl

        try:
            # data_only=True returns the cached result of formulas instead of the formula text.
            # Templates with drop-down validations make openpyxl warn ("Data Validation
            # extension is not supported"); Odoo can turn warnings into errors, so silence them.
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as e:
            raise UserError(_("The Excel file could not be opened: %s") % e)

        legend = cls._read_colour_legend(wb)
        sheets = []
        for ws in wb.worksheets:
            grid, fills = [], []
            for row in ws.iter_rows():
                grid.append([cls._to_text(c.value) for c in row])
                fills.append([cls._fill_key(c) for c in row])
            sheet = cls._build_sheet(
                ws.title, grid,
                visible=(ws.sheet_state == 'visible'),
                fills=fills, legend=legend,
            )
            sheets.append(sheet)
        return {'encoding': 'utf-8', 'delimiter': False, 'sheets': sheets}

    @classmethod
    def _parse_xls_legacy(cls, content):
        try:
            frames = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None, dtype=object)
        except Exception as e:
            raise UserError(_("The Excel file could not be opened: %s") % e)
        sheets = []
        for name, frame in frames.items():
            grid = [[cls._to_text(v) for v in row] for row in frame.itertuples(index=False)]
            sheets.append(cls._build_sheet(name, grid, visible=True, fills=None, legend=None))
        return {'encoding': 'utf-8', 'delimiter': False, 'sheets': sheets}

    # ------------------------------------------------------------------
    # Sheet building
    # ------------------------------------------------------------------

    @classmethod
    def _build_sheet(cls, name, grid, visible, fills, legend, force_data=False):
        """Classify a raw grid and turn its data area into a clean DataFrame."""
        kind = SHEET_DATA if force_data else cls._classify(name, grid)
        sheet = {
            'name': name,
            'kind': kind,
            'visible': visible,
            'header_row_index': 0,
            'df': pd.DataFrame(),
            'column_meta': {},
            'dropped_rows': 0,
        }
        if kind != SHEET_DATA:
            return sheet

        header_idx = cls._find_header_row(grid)
        if header_idx is None:
            sheet['kind'] = SHEET_EMPTY
            return sheet

        header_cells = grid[header_idx]
        body = grid[header_idx + 1:]
        width = max([len(header_cells)] + [len(r) for r in body]) if body else len(header_cells)

        def cell(row, i):
            return row[i] if i < len(row) else None

        # Keep columns that have a header or at least one value
        keep_cols, headers, seen = [], [], {}
        for i in range(width):
            raw_header = cell(header_cells, i)
            has_values = any(cell(r, i) is not None for r in body)
            if raw_header is None and not has_values:
                continue
            header = raw_header if raw_header is not None else _("Column %s") % (i + 1)
            if header in seen:
                seen[header] += 1
                header = "%s (%s)" % (header, seen[header])
            else:
                seen[header] = 1
            keep_cols.append(i)
            headers.append(header)

        if not headers:
            sheet['kind'] = SHEET_EMPTY
            return sheet

        rows, row_numbers, dropped = [], [], 0
        for offset, raw_row in enumerate(body):
            values = [cell(raw_row, i) for i in keep_cols]
            non_empty = [v for v in values if v is not None]
            if not non_empty:
                dropped += 1          # phantom row (empty or only formulas returning "")
                continue
            if cls._is_marker_row(non_empty, len(values)):
                dropped += 1          # row of 1s/xs used by some templates as a flag line
                continue
            rows.append(values)
            row_numbers.append(header_idx + offset + 2)  # 1-based spreadsheet row

        # Templates often pre-fill a default (e.g. is_company = 0) on hundreds of
        # empty rows. A row whose only values are those defaults is not real data.
        fillers = cls._filler_values(rows, len(headers))
        if fillers:
            kept_rows, kept_numbers = [], []
            for values, number in zip(rows, row_numbers):
                non_empty = [(i, v) for i, v in enumerate(values) if v is not None]
                if all(fillers.get(i) == v for i, v in non_empty):
                    dropped += 1
                    continue
                kept_rows.append(values)
                kept_numbers.append(number)
            rows, row_numbers = kept_rows, kept_numbers

        # An isolated last row made only of "1"/"x" is a template flag line
        if len(rows) >= 2:
            non_empty = [v for v in rows[-1] if v is not None]
            isolated = row_numbers[-1] > row_numbers[-2] + 1
            if isolated and len(non_empty) >= 2 and set(non_empty) <= _MARKER_VALUES:
                rows, row_numbers = rows[:-1], row_numbers[:-1]
                dropped += 1

        df = pd.DataFrame(rows, columns=headers, dtype=object)
        df.index = row_numbers

        column_meta = {}
        if fills is not None and legend:
            header_fills = fills[header_idx]
            for col_i, header in zip(keep_cols, headers):
                fill = header_fills[col_i] if col_i < len(header_fills) else None
                column_meta[header] = {'required_level': legend.get(fill)}
        else:
            for header in headers:
                column_meta[header] = {'required_level': None}

        sheet.update({
            'header_row_index': header_idx,
            'df': df,
            'column_meta': column_meta,
            'dropped_rows': dropped,
        })
        if df.empty:
            sheet['kind'] = SHEET_EMPTY
        return sheet

    @classmethod
    def _classify(cls, name, grid):
        norm_name = cls.normalize_header(name)
        first = next((v for row in grid[:3] for v in row if v is not None), '')
        norm_first = cls.normalize_header(first)

        if 'instruccion' in norm_name or 'instruction' in norm_name or norm_first.startswith('instruccion'):
            return SHEET_INSTRUCTIONS
        if 'posibles valores' in norm_name or 'possible values' in norm_name \
                or norm_first.startswith('posibles valores'):
            return SHEET_VALUES
        if not any(v is not None for row in grid for v in row):
            return SHEET_EMPTY
        return SHEET_DATA

    @classmethod
    def _find_header_row(cls, grid):
        """The header is the row with most text cells among the first rows."""
        best_idx, best_count = None, 1
        for idx, row in enumerate(grid[:15]):
            count = sum(1 for v in row if v is not None and not cls._looks_numeric(v))
            if count > best_count:
                best_idx, best_count = idx, count
        if best_idx is None:
            # Single column sheets: first non-empty row
            for idx, row in enumerate(grid[:15]):
                if any(v is not None for v in row):
                    return idx
        return best_idx

    @staticmethod
    def _filler_values(rows, width, min_count=5, ratio=0.8):
        """Columns where one value fills most cells (a template default).

        :return: {column_index: filler_value}
        """
        fillers = {}
        for i in range(width):
            counts = {}
            total = 0
            for values in rows:
                v = values[i]
                if v is not None:
                    total += 1
                    counts[v] = counts.get(v, 0) + 1
            if not counts:
                continue
            value, count = max(counts.items(), key=lambda kv: kv[1])
            if count >= min_count and count >= total * ratio:
                fillers[i] = value
        return fillers

    @staticmethod
    def _is_marker_row(non_empty, width):
        return (
            width >= 3
            and len(non_empty) >= max(3, int(width * 0.6))
            and len(set(non_empty)) == 1
            and non_empty[0] in _MARKER_VALUES
        )

    @staticmethod
    def _looks_numeric(value):
        try:
            float(str(value).replace(',', '.'))
            return True
        except ValueError:
            return False

    # ------------------------------------------------------------------
    # Colour legend (Obligatorio / Condicional / Opcional)
    # ------------------------------------------------------------------

    @classmethod
    def _read_colour_legend(cls, wb):
        """Map header fill colours to required levels using the instructions sheet."""
        legend = {}
        for ws in wb.worksheets:
            if 'instruccion' not in cls.normalize_header(ws.title) \
                    and 'instruction' not in cls.normalize_header(ws.title):
                continue
            for row in ws.iter_rows(max_row=40, max_col=12):
                for c in row:
                    word = cls.normalize_header(c.value) if isinstance(c.value, str) else ''
                    level = _LEGEND_WORDS.get(word)
                    fill = cls._fill_key(c)
                    if level and fill and fill not in legend:
                        legend[fill] = level
        return legend

    @staticmethod
    def _fill_key(cell):
        fill = getattr(cell, 'fill', None)
        if fill is None or not fill.fill_type:
            return None
        color = fill.fgColor
        ctype = getattr(color, 'type', None)
        if ctype == 'rgb':
            return ('rgb', color.rgb)
        if ctype == 'theme':
            return ('theme', color.theme, round(color.tint or 0.0, 2))
        if ctype == 'indexed':
            return ('indexed', color.indexed)
        return None

    # ------------------------------------------------------------------
    # Value conversion
    # ------------------------------------------------------------------

    @classmethod
    def _to_text(cls, value):
        """Convert any cell value to a clean string, or None when empty."""
        if value is None:
            return None
        if isinstance(value, float) and pd.isna(value):
            return None
        if isinstance(value, bool):
            return 'True' if value else 'False'
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return ('%.10f' % value).rstrip('0').rstrip('.')
        if isinstance(value, datetime.datetime):
            if value.time() == datetime.time(0, 0):
                return value.date().isoformat()
            return value.strftime('%Y-%m-%d %H:%M:%S')
        if isinstance(value, datetime.date):
            return value.isoformat()
        if isinstance(value, datetime.time):
            return value.strftime('%H:%M:%S')
        text = str(value)
        if not text.strip() or text.strip() in _EXCEL_ERRORS:
            return None
        return text

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _default_sheet(sheets):
        for predicate in (
            lambda s: s['kind'] == SHEET_DATA and s['visible'],
            lambda s: s['kind'] == SHEET_DATA,
        ):
            found = next((s for s in sheets if predicate(s)), None)
            if found:
                return found
        raise UserError(_(
            "No data sheet was found in this file. Sheets found: %s"
        ) % ', '.join('%s (%s)' % (s['name'], s['kind']) for s in sheets))

    @staticmethod
    def _sheet_summary(sheet):
        df = sheet['df']
        return {
            'name': sheet['name'],
            'kind': sheet['kind'],
            'visible': sheet['visible'],
            'row_count': len(df),
            'column_count': len(df.columns),
            'headers': [str(h) for h in df.columns],
            'dropped_rows': sheet['dropped_rows'],
        }
